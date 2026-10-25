import io

import pytest
from docx import Document as DocxDocument

from app.services.parsing import UnsupportedDocumentError, parse_document


def _docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Reward points are valid for 3 years.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Tenure", "Rate"
    table.cell(1, 0).text, table.cell(1, 1).text = "1 year", "7.10%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_paragraphs_and_tables():
    sections = parse_document("policy.docx", _docx_bytes())
    text = sections[0].text
    assert "valid for 3 years" in text
    assert "1 year | 7.10%" in text


def test_text_file_utf8_hindi():
    sections = parse_document("faq.md", "कार्ड ब्लॉक करें".encode())
    assert sections[0].text == "कार्ड ब्लॉक करें"


def test_unsupported_extension():
    with pytest.raises(UnsupportedDocumentError):
        parse_document("image.png", b"\x89PNG")


def test_empty_document_rejected():
    with pytest.raises(UnsupportedDocumentError):
        parse_document("empty.txt", b"   \n  ")
