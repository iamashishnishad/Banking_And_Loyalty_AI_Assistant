"""Extract plain text from uploaded files. Each parser returns Sections so page numbers survive
into chunk metadata (useful for citations later)."""

import io
from dataclasses import dataclass
from pathlib import PurePath

from app.services.text import normalize

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedDocumentError(ValueError):
    pass


@dataclass
class Section:
    text: str
    page: int | None = None


def _parse_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = normalize(page.extract_text() or "")
        if text:
            sections.append(Section(text=text, page=number))
    return sections


def _parse_docx(data: bytes) -> list[Section]:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables hold most fee schedules and reward-rate grids; flatten each row to "a | b | c".
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = normalize("\n\n".join(blocks))
    return [Section(text=text)] if text else []


def _parse_text(data: bytes) -> list[Section]:
    text = normalize(data.decode("utf-8", errors="replace"))
    return [Section(text=text)] if text else []


def parse_document(filename: str, data: bytes) -> list[Section]:
    ext = PurePath(filename).suffix.lower()
    parsers = {".pdf": _parse_pdf, ".docx": _parse_docx, ".txt": _parse_text, ".md": _parse_text}
    if ext not in parsers:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    try:
        sections = parsers[ext](data)
    except Exception as exc:
        # pypdf / python-docx raise their own error types for corrupt files. Retrying can't fix
        # a corrupt file, so the ingestion worker treats this as a permanent failure.
        raise UnsupportedDocumentError(f"Could not read '{filename}': {type(exc).__name__}") from exc
    if not sections:
        # Usually a scanned PDF with no text layer; OCR is out of scope for now.
        raise UnsupportedDocumentError(f"No extractable text found in '{filename}'")
    return sections


def mime_type_for(filename: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
        ".md": "text/markdown",
    }.get(PurePath(filename).suffix.lower(), "application/octet-stream")
